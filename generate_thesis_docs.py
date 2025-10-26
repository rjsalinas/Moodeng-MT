import os
import sys
from datetime import datetime

try:
    import pandas as pd
except Exception:
    pd = None

try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
except Exception:
    Document = None


PROJECT_ROOT = os.path.dirname(os.path.abspath(__file__))


def safe_read_csv(path):
    if pd is None:
        return None
    try:
        if not os.path.exists(path):
            return None
        return pd.read_csv(path)
    except Exception:
        return None


def add_heading(document, text, level=1):
    document.add_heading(text, level=level)


def add_paragraph(document, text):
    document.add_paragraph(text)


def add_bullet(document, text):
    document.add_paragraph(text, style='List Bullet')


def add_numbered(document, text):
    document.add_paragraph(text, style='List Number')


def _apply_academic_formatting(doc: "Document") -> None:
    # Page margins
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    # Base font and paragraph spacing
    normal_style = doc.styles['Normal']
    font = normal_style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    try:
        # Ensure proper font mapping for Word
        font.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    except Exception:
        pass
    pf = normal_style.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)


def _add_page_numbers(doc: "Document") -> None:
    try:
        section = doc.sections[0]
        footer = section.footer
        if footer.paragraphs:
            para = footer.paragraphs[0]
        else:
            para = footer.add_paragraph()
        para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        run = para.add_run()
        fld_begin = OxmlElement('w:fldChar')
        fld_begin.set(qn('w:fldCharType'), 'begin')
        run._r.append(fld_begin)

        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = 'PAGE'
        run._r.append(instr)

        fld_sep = OxmlElement('w:fldChar')
        fld_sep.set(qn('w:fldCharType'), 'separate')
        run._r.append(fld_sep)

        fld_end = OxmlElement('w:fldChar')
        fld_end.set(qn('w:fldCharType'), 'end')
        run._r.append(fld_end)
    except Exception:
        pass


def _add_cover_page(doc: "Document", title: str, subtitle_lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(20)
    for line in subtitle_lines:
        p = doc.add_paragraph(line)
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_page_break()


def build_methodology_doc(output_path: str) -> None:
    if Document is None:
        raise RuntimeError("python-docx is not installed. Please install python-docx and retry.")

    doc = Document()

    _apply_academic_formatting(doc)
    _add_page_numbers(doc)
    _add_cover_page(
        doc,
        'Methodology',
        [
            'Moodeng-MT: Filipino Tweet Preprocessing and Translation',
            datetime.now().strftime('%B %Y'),
        ],
    )

    # Title (document body)
    doc.add_heading('Methodology', level=0)
    p = doc.add_paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    p.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT

    # 1. Data Sources
    add_heading(doc, '1. Data Sources', level=1)
    add_bullet(doc, 'Primary parallel dataset: filipino_english_parallel_corpus.csv (text, english_translation or src, tgt)')
    add_bullet(doc, 'Enhanced dataset: full_enhanced_parallel_corpus.csv (src/tgt; optional src_enhanced/tgt_enhanced)')
    add_bullet(doc, 'Tweet normalization outputs: tweets_id_filipino_text_normalized.csv; filtered: tweets_id_filipino_text_only.csv')
    add_bullet(doc, 'Auxiliary logs: logs/normalization_log.jsonl; batch_processing.log')

    # 2. Preprocessing Pipeline
    add_heading(doc, '2. Preprocessing Pipeline', level=1)
    add_paragraph(doc, 'We apply a Filipino-aware normalization pipeline implemented in normalizer.py with English preservation and punctuation handling. Key stages and rules:')
    add_numbered(doc, 'Text cleaning and whitespace normalization')
    add_numbered(doc, 'Gibberish/keyboard-smash removal (conservative)')
    add_numbered(doc, 'Social-media artifact handling (mentions, hashtags, URLs)')
    add_numbered(doc, 'Orthographic normalization (o↔u, e↔i, etc.) and slang expansion')
    add_numbered(doc, 'Token split/merge, transpositions, and final formatting')
    add_paragraph(doc, 'Outputs preserve original terminal punctuation (?!), reducing repeats, and add a period only when none exists.')
    add_bullet(doc, 'Language filtering: Spanish confidence > 0.3 excluded; Filipino confidence > 0.1 included')
    add_bullet(doc, 'Final dataset QC: length 10–500 chars; word count 2–100; duplicate removal')
    add_bullet(doc, 'Scripts: extract_tweet_data.py → normalize_csv_tweets.py → remove_spanish_from_filipino.py → final CSV')

    # 3. CalamanCy-enhanced Preprocessing
    add_heading(doc, '3. CalamanCy-enhanced Preprocessing', level=1)
    add_paragraph(doc, 'CalamanCy (Tagalog NLP) augments the corpus with:')
    add_bullet(doc, 'Tagalog-aware tokenization and sentence boundaries')
    add_bullet(doc, 'Linguistic complexity features (POS, dependency, morphology)')
    add_bullet(doc, 'Quality validation (grammar, entities) and optional augmentation')
    add_bullet(doc, 'Enhanced columns: src_enhanced/tgt_enhanced; metadata: complexity_score, quality_score, tagalog_complexity, is_augmented')
    add_bullet(doc, 'Batch process: batch_process_calamancy.py; resumable with intermediate enhanced_batch_XXX.csv')

    # 4. Train/Validation Split
    add_heading(doc, '4. Train/Validation Split', level=1)
    add_paragraph(doc, 'Default 80/20 split, optionally guided by curriculum thresholds on complexity metrics.')

    # 5. Model and Tokenizer
    add_heading(doc, '5. Model and Tokenizer', level=1)
    add_paragraph(doc, 'Base model: facebook/mbart-large-50-many-to-many-mmt; Tokenizer: MBart50Tokenizer with src_lang=tl_XX, tgt_lang=en_XX.')

    # 6. Parameter-efficient Fine-tuning (LoRA)
    add_heading(doc, '6. Parameter-efficient Fine-tuning (LoRA)', level=1)
    add_paragraph(doc, 'LoRA adapters target attention/FFN modules for efficient training; best adapters saved to fine-tuned-mbart-tl2en-best/.')

    # 7. Curriculum and Losses
    add_heading(doc, '7. Curriculum and Losses', level=1)
    add_paragraph(doc, 'Progressive exposure from simple to complex examples using complexity thresholds. Loss mixture includes Cross-Entropy (CE), label smoothing, focal loss, and optional R-Drop.')
    cur_rows = [
        ('Phase', 'Threshold (example)', 'Loss mixture'),
        ('Simple', 'low complexity', 'CE'),
        ('Medium', '≤ mid complexity', '0.7 CE + 0.3 Label Smoothing'),
        ('Complex', '≤ high complexity', '0.5 CE + 0.3 Label Smoothing + 0.2 Focal'),
        ('Mixed', 'all data', '0.4 CE + 0.3 LS + 0.2 Focal + 0.1 R-Drop'),
    ]
    table_c = doc.add_table(rows=len(cur_rows), cols=3)
    table_c.style = 'Light Shading'
    for i, (a, b, c) in enumerate(cur_rows):
        cells = table_c.rows[i].cells
        cells[0].text = a
        cells[1].text = b
        cells[2].text = c

    # 8. Optimization and Scheduling
    add_heading(doc, '8. Optimization and Scheduling', level=1)
    add_paragraph(doc, 'AdamW optimizer; cosine schedule with warmup; gradient accumulation; mixed precision on CUDA. Typical max_seq_len=128; beam search for eval with length/repetition controls.')

    # 9. Evaluation Metrics
    add_heading(doc, '9. Evaluation Metrics', level=1)
    add_paragraph(doc, 'Validation loss and BLEU (with smoothing) on a held-out set. Sentence-level BLEU is averaged as a proxy for corpus BLEU; for formal reporting, compute corpus-level BLEU on the full validation/test set.')

    # 10. Inference
    add_heading(doc, '10. Inference', level=1)
    add_paragraph(doc, 'Use translate_with_model.py to load the best adapter and translate Filipino text or batch CSV columns.')
    add_bullet(doc, 'Single text: python translate_with_model.py --text "kamusta ka?"')
    add_bullet(doc, 'Batch CSV: python translate_with_model.py --input_csv file.csv --src_col src --out_csv out.csv')
    add_bullet(doc, 'Adapter path override: --model_dir fine-tuned-mbart-tl2en-best')

    # 11. Reproducibility
    add_heading(doc, '11. Reproducibility', level=1)
    add_bullet(doc, 'Record seeds, learning rate schedule, and LoRA hyperparameters')
    add_bullet(doc, 'Cache and version full_enhanced_parallel_corpus.csv; preserve batch_processing.log')
    add_bullet(doc, 'Environment: Python + PyTorch + Transformers + PEFT; CalamanCy/spaCy versions noted')

    # References
    add_heading(doc, 'References', level=1)
    refs = [
        'Tang, Y., et al. (2020). Multilingual Translation with mBART-50.',
        'Hu, E. J., et al. (2022). LoRA: Low-Rank Adaptation of Large Language Models.',
        'Honnibal, M., et al. spaCy: Industrial-Strength NLP.',
        'CalamanCy: Tagalog NLP toolkit (GitHub project).',
    ]
    for r in refs:
        add_bullet(doc, r)

    try:
        doc.save(output_path)
    except PermissionError:
        alt = output_path.replace('.docx', f"_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")
        doc.save(alt)
        print(f"Target locked, saved methodology as: {alt}")


def build_results_doc(output_path: str) -> None:
    if Document is None:
        raise RuntimeError("python-docx is not installed. Please install python-docx and retry.")

    doc = Document()
    _apply_academic_formatting(doc)
    _add_page_numbers(doc)
    _add_cover_page(
        doc,
        'Results',
        [
            'Moodeng-MT: Filipino Tweet Preprocessing and Translation',
            datetime.now().strftime('%B %Y'),
        ],
    )
    doc.add_heading('Results', level=0)
    p = doc.add_paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    p.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT

    # Section: Preprocessing Results
    add_heading(doc, '1. Preprocessing Results', level=1)
    tweets_norm_path = os.path.join(PROJECT_ROOT, 'tweets_id_filipino_text_normalized.csv')
    tweets_fil_only_path = os.path.join(PROJECT_ROOT, 'tweets_id_filipino_text_only.csv')

    df_norm = safe_read_csv(tweets_norm_path)
    df_fil = safe_read_csv(tweets_fil_only_path)

    # Table: Preprocessing summary
    prep_rows = []
    if df_norm is not None:
        non_empty = 0
        if df_norm is not None and 'preprocessed_text' in df_norm.columns:
            try:
                non_empty = int(df_norm['preprocessed_text'].dropna().shape[0])
            except Exception:
                non_empty = 0
        prep_rows.append(('Normalized tweets (rows)', f"{len(df_norm):,}"))
        prep_rows.append(('Non-empty normalized texts', f"{non_empty:,}"))
        # Derived stats
        try:
            texts = df_norm['preprocessed_text'].dropna().astype(str)
            avg_len = int(texts.str.len().mean()) if not texts.empty else 0
            avg_words = float(texts.str.split().map(len).mean()) if not texts.empty else 0.0
            prep_rows.append(('Average text length (chars)', f"{avg_len}"))
            prep_rows.append(('Average word count', f"{avg_words:.1f}"))
        except Exception:
            pass
    else:
        prep_rows.append(('Normalized tweets file', 'Not found'))

    if df_fil is not None:
        prep_rows.append(('Filtered Filipino/Taglish tweets', f"{len(df_fil):,}"))
    else:
        prep_rows.append(('Filtered Filipino tweets file', 'Not found'))

    table = doc.add_table(rows=max(1, len(prep_rows))+1, cols=2)
    table.style = 'Light Shading'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Metric'
    hdr_cells[1].text = 'Value'
    for r, (k, v) in enumerate(prep_rows, start=1):
        cells = table.rows[r].cells
        cells[0].text = k
        cells[1].text = v

    add_paragraph(doc, 'Normalization preserved English segments and original terminal punctuation while removing repeated marks and adding periods only when needed.')

    # Section: CalamanCy-enhanced Corpus
    add_heading(doc, '2. CalamanCy-enhanced Corpus', level=1)
    enhanced_path = os.path.join(PROJECT_ROOT, 'full_enhanced_parallel_corpus.csv')
    df_enh = safe_read_csv(enhanced_path)
    enh_rows = []
    if df_enh is not None:
        enh_rows.append(('Enhanced pairs (rows)', f"{len(df_enh):,}"))
        present_cols = [c for c in ['src', 'tgt', 'src_enhanced', 'tgt_enhanced', 'complexity_score', 'quality_score', 'tagalog_complexity', 'is_augmented'] if c in df_enh.columns]
        enh_rows.append(('Columns present', ', '.join(present_cols) if present_cols else '(none detected)'))
        if 'is_augmented' in df_enh.columns:
            try:
                aug_count = int(df_enh['is_augmented'].fillna(False).astype(bool).sum())
                enh_rows.append(('Augmented rows', f"{aug_count:,}"))
            except Exception:
                pass
        # Summary of complexity/quality
        try:
            if 'complexity_score' in df_enh.columns:
                enh_rows.append(('Avg complexity_score', f"{float(df_enh['complexity_score'].dropna().mean()):.3f}"))
            if 'quality_score' in df_enh.columns:
                enh_rows.append(('Avg quality_score', f"{float(df_enh['quality_score'].dropna().mean()):.3f}"))
        except Exception:
            pass
    else:
        enh_rows.append(('Enhanced corpus', 'Not found'))

    table2 = doc.add_table(rows=max(1, len(enh_rows))+1, cols=2)
    table2.style = 'Light Shading'
    hdr2 = table2.rows[0].cells
    hdr2[0].text = 'Metric'
    hdr2[1].text = 'Value'
    for r, (k, v) in enumerate(enh_rows, start=1):
        cells = table2.rows[r].cells
        cells[0].text = k
        cells[1].text = v

    # Section: Model Training Summary
    add_heading(doc, '3. Model Training Summary', level=1)
    best_dir = os.path.join(PROJECT_ROOT, 'fine-tuned-mbart-tl2en-best')
    checkpoints_dir = os.path.join(PROJECT_ROOT, 'fine-tuned-mbart-tl2en')
    art_rows = [
        ('Best adapter directory present', 'Yes' if os.path.isdir(best_dir) else 'No'),
        ('Checkpoints directory present', 'Yes' if os.path.isdir(checkpoints_dir) else 'No'),
    ]
    table3 = doc.add_table(rows=len(art_rows)+1, cols=2)
    table3.style = 'Light Shading'
    h = table3.rows[0].cells
    h[0].text = 'Artifact'
    h[1].text = 'Status'
    for r, (k, v) in enumerate(art_rows, start=1):
        cells = table3.rows[r].cells
        cells[0].text = k
        cells[1].text = v

    # Section: Evaluation Metrics
    add_heading(doc, '4. Evaluation Metrics', level=1)
    add_paragraph(doc, 'If available, validation loss and BLEU are reported from training logs. For full corpus BLEU, increase evaluation sample size or run a dedicated evaluation script.')

    # Try to scan a simple BLEU/val loss from any text logs under training_logs/
    logs_root = os.path.join(PROJECT_ROOT, 'training_logs')
    extracted = []
    if os.path.isdir(logs_root):
        for root, _, files in os.walk(logs_root):
            for name in files:
                if name.lower().endswith(('.txt', '.log')):
                    try:
                        with open(os.path.join(root, name), 'r', encoding='utf-8', errors='ignore') as fh:
                            content = fh.read()
                        # Very light heuristic extraction
                        # Look for patterns like "bleu=..", "avg_bleu=..", "final_bleu=..", "val_loss=..", "best_val_loss=..", "final_val_loss=.."
                        import re
                        bleu_matches = []
                        for pat in [r"avg_bleu\s*[:=]\s*([0-9]+\.?[0-9]*)", r"bleu\s*[:=]\s*([0-9]+\.?[0-9]*)", r"final_bleu\s*[:=]\s*([0-9]+\.?[0-9]*)", r"best_bleu\s*[:=]\s*([0-9]+\.?[0-9]*)"]:
                            bleu_matches += re.findall(pat, content, flags=re.IGNORECASE)
                        vloss_matches = []
                        for pat in [r"val(?:idation)?[_\-\s]*loss\s*[:=]\s*([0-9]+\.?[0-9]*)", r"best_val_loss\s*[:=]\s*([0-9]+\.?[0-9]*)", r"final_val_loss\s*[:=]\s*([0-9]+\.?[0-9]*)"]:
                            vloss_matches += re.findall(pat, content, flags=re.IGNORECASE)
                        if bleu_matches or vloss_matches:
                            extracted.append((name, bleu_matches[-1] if bleu_matches else None, vloss_matches[-1] if vloss_matches else None))
                    except Exception:
                        pass
    if extracted:
        add_paragraph(doc, 'Detected metrics from logs (last entries):')
        met_table = doc.add_table(rows=min(5, len(extracted))+1, cols=3)
        met_table.style = 'Light Shading'
        mh = met_table.rows[0].cells
        mh[0].text = 'File'
        mh[1].text = 'val_loss'
        mh[2].text = 'avg_bleu'
        for i, (fname, bleu, vloss) in enumerate(extracted[-5:], start=1):
            cells = met_table.rows[i].cells
            cells[0].text = fname
            cells[1].text = vloss if vloss is not None else '-'
            cells[2].text = bleu if bleu is not None else '-'
    else:
        add_paragraph(doc, 'No metrics detected in training_logs/.')

    # References
    add_heading(doc, 'References', level=1)
    refs = [
        'Tang, Y., et al. (2020). Multilingual Translation with mBART-50.',
        'Hu, E. J., et al. (2022). LoRA: Low-Rank Adaptation of Large Language Models.',
        'Honnibal, M., et al. spaCy: Industrial-Strength NLP.',
        'CalamanCy: Tagalog NLP toolkit (GitHub project).',
    ]
    for r in refs:
        add_bullet(doc, r)

    try:
        doc.save(output_path)
    except PermissionError:
        alt = output_path.replace('.docx', f"_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")
        doc.save(alt)
        print(f"Target locked, saved results as: {alt}")


def main():
    out_method = os.path.join(PROJECT_ROOT, 'Thesis_Methodology.docx')
    out_results = os.path.join(PROJECT_ROOT, 'Thesis_Results.docx')

    build_methodology_doc(out_method)
    build_results_doc(out_results)

    print(f"\n✓ Documents generated:\n- {out_method}\n- {out_results}")


if __name__ == '__main__':
    main()


